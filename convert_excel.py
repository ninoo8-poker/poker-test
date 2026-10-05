"""Convertit l'Excel de ranges en data/ranges.json.

Usage : python convert_excel.py range_KT_100BB.xlsx

Format attendu (identique au fichier d'origine) :
- 1 feuille par position, titre en A1
- grille 13x13 en A2:M14 (paires sur la diagonale, suited au-dessus, offsuit en dessous)
- couleur de fond : rouge FF0000 = raise, vert 00B050 = call, autre (gris) = fold
- taille du raise en C17 (ex. « RAISE 2,1 »)
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

RANKS = "AKQJT98765432"


def expected(i, j):
    if i == j:
        return RANKS[i] * 2
    return RANKS[i] + RANKS[j] + "s" if i < j else RANKS[j] + RANKS[i] + "o"


def convert(xlsx_path, out_path="data/ranges.json"):
    wb = load_workbook(xlsx_path)
    out = {}
    for ws in wb.worksheets:
        pos = str(ws["A1"].value).strip()
        size = str(ws["C17"].value).upper().replace("RAISE", "").strip()
        actions = {}
        for i in range(13):
            for j in range(13):
                cell = ws.cell(row=2 + i, column=1 + j)
                if cell.value != expected(i, j):
                    raise ValueError(f"{ws.title}!{cell.coordinate}: {cell.value!r} ≠ {expected(i, j)!r}")
                f = cell.fill.fgColor if cell.fill.fill_type else None
                rgb = f.rgb if f is not None and f.type == "rgb" else None
                actions[cell.value] = {"FFFF0000": "raise", "FF00B050": "call"}.get(rgb, "fold")
        out[pos] = {"raise_size": size, "actions": actions}
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    Path(out_path).write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(out)} positions écrites dans {out_path}")


if __name__ == "__main__":
    convert(sys.argv[1] if len(sys.argv) > 1 else "range_KT_100BB.xlsx")
